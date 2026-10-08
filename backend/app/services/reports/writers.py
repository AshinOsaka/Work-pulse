"""Report files: CSV, Excel (XLSX) and PDF from the same typed rows.

* Durations are hours with two decimals in CSV/XLSX (sortable, summable) and "6h 42m" in PDF.
* Cells that a spreadsheet would treat as a formula (=, +, -, @, tab, CR) are prefixed with an apostrophe:
  task titles, application names and the like come from users and must never execute (CSV/formula injection).
* PDFs use a Unicode TrueType font when one is available so every name renders; without one, text is reduced
  to Latin-1 and the report says so.
"""

from __future__ import annotations

import csv
import io
import os
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.reports.builders import Column, ReportData

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
FONT_CANDIDATES = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "C:/Windows/Fonts/segoeui.ttf",
    "C:/Windows/Fonts/arial.ttf",
    "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
)


@dataclass(frozen=True, slots=True)
class ReportMeta:
    title: str
    subtitle: str
    lines: list[str]
    notes: list[str]


def safe_text(value: str) -> str:
    return f"'{value}" if value.startswith(FORMULA_PREFIXES) else value


def _hours(seconds: float | None) -> float | None:
    return None if seconds is None else round(seconds / 3600, 2)


def _duration(seconds: float | None) -> str:
    if seconds is None:
        return "—"
    minutes = round(seconds / 60)
    hours, rest = divmod(minutes, 60)
    return f"{hours}h {rest:02d}m" if hours else f"{minutes}m"


def header(column: Column) -> str:
    return (
        f"{column.label} (h)"
        if column.kind == "duration"
        else f"{column.label} (%)"
        if column.kind == "percent"
        else column.label
    )


def _plain(column: Column, value: Any) -> Any:
    """Value for CSV / XLSX."""
    if value is None or value == "":
        return ""
    if column.kind == "duration":
        return _hours(value)
    if column.kind == "time" and isinstance(value, datetime):
        return value.strftime("%H:%M")
    if column.kind == "datetime" and isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if column.kind == "date" and isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return safe_text(value)
    return value


def _display(column: Column, value: Any) -> str:
    """Value for PDF."""
    if value is None or value == "":
        return "" if column.kind == "text" else "—"
    if column.kind == "duration":
        return _duration(value)
    if column.kind == "percent":
        return f"{value}%"
    if column.kind == "time" and isinstance(value, datetime):
        return value.strftime("%H:%M")
    if column.kind == "datetime" and isinstance(value, datetime):
        return value.strftime("%d %b %Y %H:%M")
    if column.kind == "date" and isinstance(value, date):
        return value.strftime("%d %b %Y")
    return str(value)


def to_csv(data: ReportData, meta: ReportMeta) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([header(c) for c in data.columns])
    for row in data.rows:
        writer.writerow([_plain(c, row.get(c.key)) for c in data.columns])
    # UTF-8 with BOM so Excel opens non-ASCII names correctly.
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def to_xlsx(data: ReportData, meta: ReportMeta) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = data.title[:31]
    bold = Font(bold=True)
    sheet.append([header(c) for c in data.columns])
    for cell in sheet[1]:
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="EEF0F5")
        cell.alignment = Alignment(vertical="center")
    for row in data.rows:
        sheet.append([_plain(c, row.get(c.key)) for c in data.columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(data.columns, start=1):
        widest = max(
            [len(header(column)), *(len(str(_plain(column, r.get(column.key)))) for r in data.rows[:500])]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(48, max(9, widest + 2))
        if column.kind == "duration":
            for cell in sheet[get_column_letter(index)][1:]:
                cell.number_format = "0.00"
    about = book.create_sheet("About this report")
    for line in [meta.title, meta.subtitle, "", *meta.lines, "", *meta.notes]:
        about.append([safe_text(line)])
    about["A1"].font = Font(bold=True, size=14)
    about.column_dimensions["A"].width = 120
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _font_path(configured: str) -> str | None:
    for candidate in (configured, *FONT_CANDIDATES):
        if candidate and os.path.isfile(candidate):
            return candidate
    return None


class _Pdf(FPDF):  # type: ignore[misc]  # fpdf2 ships without type stubs
    def __init__(self, title: str, family: str) -> None:
        super().__init__(orientation="L", unit="mm", format="A4")
        self.report_title = title
        self.family = family

    def footer(self) -> None:
        self.set_y(-10)
        self.set_font(self.family, size=7)
        self.set_text_color(120, 120, 120)
        self.cell(0, 5, f"{self.report_title} · WorkPulse · page {self.page_no()} of {{nb}}", align="R")


def to_pdf(data: ReportData, meta: ReportMeta, *, max_rows: int, font: str = "") -> tuple[bytes, list[str]]:
    """Returns the file and any notes about it (row limit, font fallback)."""
    extra: list[str] = []
    path = _font_path(font)
    pdf = _Pdf(meta.title, "Body" if path else "Helvetica")
    unicode = path is not None
    if path:
        pdf.add_font("Body", "", path)
        pdf.add_font("Body", "B", path)  # same face; weight is conveyed by size and colour
    else:
        extra.append("No Unicode font was available, so characters outside Latin-1 are replaced in this PDF.")

    def fix(text: str) -> str:
        if unicode:
            return text
        for a, b in (("—", "-"), ("–", "-"), ("×", "x"), ("÷", "/"), ("•", "-")):  # noqa: RUF001
            text = text.replace(a, b)
        return text.encode("latin-1", "replace").decode("latin-1")

    rows = data.rows
    if len(rows) > max_rows:
        extra.append(
            f"This PDF shows the first {max_rows:,} of {len(rows):,} rows. CSV and Excel include every row."
        )
        rows = rows[:max_rows]
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.set_margins(10, 10, 10)
    pdf.add_page()
    pdf.set_font(pdf.family, "B", 15)
    pdf.cell(0, 8, fix(meta.title), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(pdf.family, size=9)
    pdf.set_text_color(90, 90, 90)
    pdf.cell(0, 5, fix(meta.subtitle), new_x="LMARGIN", new_y="NEXT")
    for line in meta.lines:
        pdf.cell(0, 4.5, fix(line), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(1)
    for note in [*meta.notes, *extra]:
        pdf.multi_cell(0, 4.2, fix(f"• {note}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    pdf.set_text_color(0, 0, 0)
    if not rows:
        pdf.set_font(pdf.family, size=10)
        pdf.cell(0, 8, "No rows for this selection.", new_x="LMARGIN", new_y="NEXT")
        return bytes(pdf.output()), extra

    widths = _widths(data, rows, pdf.epw)
    pdf.set_font(pdf.family, size=7)
    numeric = {"int", "duration", "percent"}
    with pdf.table(
        col_widths=widths,
        text_align=tuple("RIGHT" if c.kind in numeric else "LEFT" for c in data.columns),
        line_height=4.2,
        headings_style=_heading_style(pdf.family),
        borders_layout="HORIZONTAL_LINES",
        cell_fill_color=(246, 247, 250),
        cell_fill_mode="ROWS",
        repeat_headings=1,
    ) as table:
        head = table.row()
        for c in data.columns:
            head.cell(fix(c.label))
        for row in rows:
            r = table.row()
            for c in data.columns:
                r.cell(fix(_display(c, row.get(c.key))))
    return bytes(pdf.output()), extra


def _heading_style(family: str) -> Any:
    from fpdf.fonts import FontFace

    return FontFace(family=family, emphasis="BOLD", color=(60, 60, 70), fill_color=(230, 232, 240))


def _widths(data: ReportData, rows: list[dict[str, Any]], total: float) -> tuple[float, ...]:
    sample = rows[:200]
    raw = []
    for c in data.columns:
        longest = max([len(c.label), *(len(_display(c, r.get(c.key))) for r in sample)])
        raw.append(min(36, max(5, longest)))
    scale = total / sum(raw)
    return tuple(w * scale for w in raw)


FORMATS = {
    "csv": ("text/csv; charset=utf-8", "csv"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "pdf": ("application/pdf", "pdf"),
}


def font_available(configured: str) -> bool:
    return _font_path(configured) is not None


__all__ = ["FORMATS", "ReportMeta", "font_available", "safe_text", "to_csv", "to_pdf", "to_xlsx"]
