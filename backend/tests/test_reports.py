"""Phase 13: reports — permissions, scope, preview, asynchronous generation, formats, secure download, audit."""

from __future__ import annotations

import csv
import io
import time
from datetime import UTC, datetime
from typing import Any
from urllib.parse import parse_qs, urlparse

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.org_helpers import Workspace, auth
from tests.test_agent_api import enrol_member
from tests.test_productivity_api import DAY, workday

D = DAY.date().isoformat()
TODAY = datetime.now(UTC).date().isoformat()


def body(report_type: str, fmt: str = "csv", **extra: Any) -> dict[str, Any]:
    return {"report_type": report_type, "format": fmt, "period": "daily", "start": D, "end": D, **extra}


def wait_ready(
    ws: Workspace, report_id: str, token: str | None = None, timeout: float = 30
) -> dict[str, Any]:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        job = ws.get(f"/reports/{report_id}", token=token).json()
        if job["status"] in ("ready", "failed"):
            return job  # type: ignore[no-any-return]
        time.sleep(0.2)
    raise AssertionError("report did not finish in time")


def download(client: TestClient, job: dict[str, Any]) -> Any:
    return client.get(job["download_url"])


def test_catalogue_and_permissions(ws: Workspace) -> None:
    lead = ws.member("Tia Lead", role="TEAM_LEAD")
    worker = ws.member("Eli Employee")
    types = {t["key"]: t for t in ws.get("/reports/types").json()}
    assert len(types) == 10 and all(t["allowed"] for t in types.values())
    assert set(types["tasks"]["filters"]) >= {"project", "employee", "department", "team", "role"}

    lead_types = {t["key"]: t["allowed"] for t in ws.get("/reports/types", token=lead.token).json()}
    assert lead_types["attendance"] and not lead_types["live_sessions"]  # team leads can't view live sessions
    assert ws.post("/reports/preview", body("live_sessions"), token=lead.token).status_code == 403
    denied = ws.post("/reports", body("attendance"), token=lead.token)  # REPORT_VIEW but not REPORT_EXPORT
    assert denied.status_code == 403
    assert ws.get("/reports/types", token=worker.token).status_code == 403

    too_long = ws.post("/reports", {**body("attendance"), "start": "2024-01-01", "end": "2026-01-01"})
    assert too_long.status_code == 400 and too_long.json()["error"]["code"] == "range_too_long"
    assert (
        ws.post(
            "/reports/preview", {**body("attendance"), "start": "2026-01-01", "end": "2026-03-01"}
        ).status_code
        == 400
    )
    backwards = ws.post("/reports", {**body("attendance"), "start": D, "end": "2020-01-01"})
    assert backwards.json()["error"]["code"] == "invalid_range"


def test_preview_is_scoped_and_correct(ws: Workspace) -> None:
    manager = ws.member("Max Manager", role="MANAGER")
    _report_member, agent = enrol_member(ws, "Rhea Report", manager_employee_id=manager.employee_id)
    other, other_agent = enrol_member(ws, "Otto Other")
    workday(agent)
    workday(other_agent)

    mine = ws.post("/reports/preview", body("attendance"), token=manager.token).json()
    names = {r["employee"] for r in mine["rows"]}
    assert names == {"Rhea Report", "Max Manager"} and mine["people"] == 2  # the reporting line (and self)
    row = next(r for r in mine["rows"] if r["employee"] == "Rhea Report")
    assert next(r for r in mine["rows"] if r["employee"] == "Max Manager")["status"] == "No work session"
    assert row["status"] == "Worked" and row["first_start"] == "09:00" and row["last_end"] == "15:00"
    assert row["work_seconds"] == "5h 00m" and row["sessions"] == "2"
    assert any("not mean absence" in n for n in mine["notes"])

    everyone = ws.post("/reports/preview", body("work_hours")).json()
    assert {"Rhea Report", "Otto Other"} <= {r["employee"] for r in everyone["rows"]}
    filtered = ws.post(
        "/reports/preview", body("applications", filters={"employee_ids": [other.employee_id]})
    ).json()
    assert {r["employee"] for r in filtered["rows"]} == {"Otto Other"}
    assert {r["key"] for r in filtered["rows"]} >= {"code.exe", "notion.exe"}
    assert ws.post("/reports/preview", body("productivity")).json()["total_rows"] >= 2


def test_generation_formats_download_and_audit(ws: Workspace, client: TestClient, sync_db: Any) -> None:
    _member, agent = enrol_member(ws, "Fay Format")
    workday(agent)
    jobs = {}
    for fmt in ("csv", "xlsx", "pdf"):
        created = ws.post("/reports", body("work_hours", fmt))
        assert created.status_code == 202, created.text
        assert created.json()["status"] == "queued"
        jobs[fmt] = wait_ready(ws, created.json()["id"])
    for fmt, job in jobs.items():
        assert job["status"] == "ready", job
        assert job["row_count"] >= 1 and job["filename"].endswith(f".{fmt}") and job["expires_at"]
        response = download(client, job)
        assert response.status_code == 200
        assert response.headers["content-disposition"].startswith("attachment;")
        assert response.headers["cache-control"] == "private, no-store"
        if fmt == "csv":
            rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
            assert rows[0][:2] == ["Date", "Employee"] and "Work (h)" in rows[0]
            fay = next(r for r in rows if "Fay Format" in r)
            assert fay[rows[0].index("Work (h)")] == "5.0"
        elif fmt == "xlsx":
            book = load_workbook(io.BytesIO(response.content))
            assert book.sheetnames == ["Work hours", "About this report"]
            assert book["Work hours"]["A1"].value == "Date"
        else:
            assert response.content.startswith(b"%PDF")

    # The signature is bound to the requester and expires; tampering is refused.
    url = urlparse(jobs["csv"]["download_url"])
    query = {k: v[0] for k, v in parse_qs(url.query).items()}
    other = ws.member("Nosy Neighbour", role="MANAGER")
    forged = client.get(url.path, params={**query, "u": other.user_id})
    assert forged.status_code == 403
    assert client.get(url.path, params={**query, "e": str(int(query["e"]) + 60)}).status_code == 403
    assert ws.get(f"/reports/{jobs['csv']['id']}", token=other.token).status_code == 404  # not their report
    assert ws.get("/reports", token=other.token).json() == []

    actions = [
        e["action"]
        for e in sync_db["audit_logs"].find({"target_id": jobs["csv"]["id"]}).sort("created_at", 1)
    ]
    assert actions == ["report.requested", "report.generated", "report.downloaded"]
    stored = sync_db["report_jobs"].find_one({"filename": jobs["csv"]["filename"]})
    assert stored["object_key"].startswith("reports/") and b"Fay Format" not in str(stored).encode()

    deleted = ws.client.delete(f"/api/reports/{jobs['pdf']['id']}", headers=auth(ws.admin.token))
    assert deleted.status_code == 204
    assert download(client, jobs["pdf"]).status_code == 404


def test_exports_neutralise_spreadsheet_formulas(ws: Workspace, client: TestClient) -> None:
    project = ws.post("/projects", {"name": "Formula", "key": "FX", "member_ids": []}).json()
    ws.post("/tasks", {"project_id": project["id"], "title": '=HYPERLINK("http://evil.example","x")'})
    job = wait_ready(ws, ws.post("/reports", body("tasks", "csv", start=TODAY, end=TODAY)).json()["id"])
    assert job["status"] == "ready", job
    text = download(client, job).content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    title = rows[1][rows[0].index("Title")]
    assert title.startswith("'=")  # shown as text, never evaluated
    xlsx = wait_ready(ws, ws.post("/reports", body("tasks", "xlsx", start=TODAY, end=TODAY)).json()["id"])
    sheet = load_workbook(io.BytesIO(download(client, xlsx).content))["Tasks"]
    assert str(sheet["B2"].value).startswith("'=")


def test_projects_and_live_session_reports(ws: Workspace) -> None:
    member = ws.member("Pia Project")
    project = ws.post(
        "/projects", {"name": "Reported", "key": "REP", "member_ids": [member.employee_id]}
    ).json()
    t = ws.post(
        "/tasks", {"project_id": project["id"], "title": "Done", "assignee_ids": [member.employee_id]}
    ).json()
    ws.patch(f"/tasks/{t['id']}", {"status": "COMPLETED"})
    ws.post(f"/tasks/{t['id']}/time", {"minutes": 45, "day": D}, token=member.token)
    preview = ws.post("/reports/preview", body("projects", filters={"project_id": project["id"]})).json()
    (row,) = preview["rows"]
    assert row["project"] == "Reported" and row["completed"] == "1" and row["progress"] == "100%"
    assert row["time_in_period"] == "45m"
    live = ws.post("/reports/preview", body("live_sessions")).json()
    assert live["total_rows"] == 0 and any("never recorded" in n for n in live["notes"])


def test_pdf_falls_back_without_a_unicode_font_and_caps_rows(monkeypatch: Any) -> None:
    from app.services.reports import writers
    from app.services.reports.builders import Column, ReportData

    monkeypatch.setattr(writers, "FONT_CANDIDATES", ())
    data = ReportData(
        "Names",
        [Column("name", "Name"), Column("seconds", "Time", "duration")],
        [{"name": f"Zoë — 日本 {i}", "seconds": 3600 * i} for i in range(150)],
    )
    meta = writers.ReportMeta(
        "Names report",
        "1 Jan – 2 Jan",  # noqa: RUF001 - non-Latin-1 on purpose
        ["Generated by test"],
        ["A note · with symbols"],
    )
    content, notes = writers.to_pdf(data, meta, max_rows=100, font="")
    assert content.startswith(b"%PDF")
    assert any("Latin-1" in n for n in notes) and any("first 100 of 150" in n for n in notes)


def test_worker_rechecks_permissions_before_generating(ws: Workspace, sync_db: Any) -> None:
    from bson import ObjectId

    employee = ws.member("Ned Noexport")
    company_id = sync_db["users"].find_one({"_id": ObjectId(employee.user_id)})["company_id"]
    now = datetime.now(UTC)
    job_id = (
        sync_db["report_jobs"]
        .insert_one(
            {
                "company_id": company_id,
                "requested_by": ObjectId(employee.user_id),
                "requested_by_name": "Ned Noexport",
                "report_type": "attendance",
                "format": "csv",
                "period": "daily",
                "start": D,
                "end": D,
                "filters": {},
                "status": "queued",
                "progress": 0,
                "notes": [],
                "downloads": 0,
                "created_at": now,
                "updated_at": now,
            }
        )
        .inserted_id
    )
    deadline = time.monotonic() + 20
    while time.monotonic() < deadline:
        doc = sync_db["report_jobs"].find_one({"_id": job_id})
        if doc["status"] == "failed":
            break
        time.sleep(0.2)
    assert doc["status"] == "failed" and "no longer have permission" in doc["error"]
    assert doc.get("object_key") is None
    assert sync_db["audit_logs"].find_one({"target_id": str(job_id), "action": "report.failed"}) is not None
